from pathlib import Path

from openpyxl import Workbook, load_workbook

from backcheck.workbook_writer import write_backcheck_workbook


def test_write_backcheck_workbook_removes_sample_sheets_and_fills_values(tmp_path: Path):
    template = tmp_path / "template.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Template"
    ws["A1"] = "title"
    ws["B4"] = "old website"
    wb.create_sheet("Sample 2")
    wb.save(template)

    output = tmp_path / "out.xlsx"
    write_backcheck_workbook(
        template,
        output,
        {"company_name": "Acme Dental", "country": "Thailand", "website": "https://acme.example"},
    )

    saved = load_workbook(output)
    assert saved.sheetnames == ["Acme Dental"]
    ws = saved["Acme Dental"]
    assert ws["A1"].value == "Thailand+Acme Dental"
    assert ws["B4"].value == "https://acme.example"
    assert ws["B7"].value is None
